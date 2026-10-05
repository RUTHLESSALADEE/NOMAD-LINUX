import pytest

from test_ipam_server import key_for, offline_store, online_again, plan, server, team_store  # noqa: F401 (fixture)

from nomad.ipam.client import ServerUnreachable
from nomad.ipam.server import ADMIN
from nomad.ipam.store import IpamError, IpamStore
from nomad.ipam.vlan_team import TeamVlanStore
from nomad.ipam.vlans import PLANNED, RESERVED, VlanStore, domain_history, name_problem, range_for, \
    suggested_links, vlan_history, vlan_named_in
from nomad.ipam.workbook import export_workbook


@pytest.fixture
def local(tmp_path):
    store = IpamStore(str(tmp_path / "ipam.db"), user="tester")
    yield store
    store.close()


def test_domains_and_vlans(local):
    network = local.add_network("11AB SIPR")
    local.add_subnet(network.id, "10.6.0.0/16", "Vlan 6")
    vlans = VlanStore(local)
    domain = vlans.add_domain("SIPR switches", network.id, "SIPR-VTP", ranges=[{"first": 100, "last": 199,
                                                                              "name": "Users"}])
    with pytest.raises(IpamError, match="already a VLAN domain"):
        vlans.add_domain("sipr SWITCHES")
    vlan = vlans.set_vlan(domain.id, 6, "MGMT", subnets=["10.6.0.0 255.255.0.0"])
    assert (vlan.vlan, vlan.name, vlan.subnets, vlan.status) == (6, "MGMT", ["10.6.0.0/16"], "active")
    assert vlans.vlan_with_subnet(domain.id, "10.6.0.0/16").vlan == 6
    with pytest.raises(IpamError, match="already in VLAN 6"):
        vlans.set_vlan(domain.id, 7, "OTHER", subnets=["10.6.0.0/16"])  # A subnet is in one VLAN of a domain
    assert vlans.set_vlan(domain.id, 6, "MGMT-2", subnets=["10.6.0.0/16"]).version == 2
    vlans.set_vlan(domain.id, 100, "USERS-1", status=PLANNED)
    vlans.set_vlan(domain.id, 101, status=RESERVED)
    assert [item.vlan for item in vlans.vlans(domain.id)] == [6, 100, 101]
    assert vlans.next_free(domain.id) == 1
    assert vlans.next_free(domain.id, 100, 199) == 102
    assert range_for(vlans.domain(domain.id), 150)["name"] == "Users"
    for bad in (0, 4095, "x"):
        with pytest.raises(IpamError):
            vlans.set_vlan(domain.id, bad)
    vlans.delete_vlan(domain.id, 100)
    assert vlans.deleted_numbers(domain.id) == {100}
    assert [domain for domain, _ in vlans.search("users")] == []  # Deleted
    assert [vlan.vlan for _, vlan in vlans.search("10.6.")] == [6]
    vlans.delete_domain(domain.id)
    assert vlans.domains() == [] and vlans.vlans(domain.id) == []


def test_vlans_leave_the_network_and_its_export_alone(local, tmp_path):
    network = local.add_network("11AB SIPR")
    local.add_subnet(network.id, "10.6.0.0/16", "Vlan 6", "10.6.0.1")
    local.set_address(network.id, "10.6.0.5", name="sw1")
    before = [(subnet.cidr, subnet.name, subnet.version) for subnet in local.subnets(network.id)]
    export_workbook(tmp_path / "before.xlsx", [(local, local.network(network.id))])
    vlans = VlanStore(local)
    domain = vlans.add_domain("SIPR", network.id)
    vlans.set_vlan(domain.id, 6, "MGMT", subnets=["10.6.0.0/16"])
    assert [(subnet.cidr, subnet.name, subnet.version) for subnet in local.subnets(network.id)] == before
    export_workbook(tmp_path / "after.xlsx", [(local, local.network(network.id))])
    from openpyxl import load_workbook
    rows = [[list(row) for row in load_workbook(tmp_path / name).active.iter_rows(values_only=True)]
            for name in ("before.xlsx", "after.xlsx")]
    assert rows[0] == rows[1]


def test_links_suggested_from_subnet_names_and_details(local):
    network = local.add_network("11AB")
    local.add_subnet(network.id, "172.28.0.0/16", "Vlan 6")
    local.add_subnet(network.id, "10.10.10.0/24", "68890 MGT", fields={"Colorless": "Vlan 10"})
    local.add_subnet(network.id, "10.20.0.0/24", "Printers")
    vlans = VlanStore(local)
    domain = vlans.add_domain("11AB", network.id)
    found = suggested_links(local, vlans, domain)
    assert [(number, cidr) for number, cidr, _ in found] == [(10, "10.10.10.0/24"), (6, "172.28.0.0/16")]
    assert "Colorless says Vlan 10" in found[0][2]
    vlans.set_vlan(domain.id, 6, subnets=["172.28.0.0/16"])
    assert [number for number, _, _ in suggested_links(local, vlans, domain)] == [10]
    assert vlan_named_in("VLAN-12 users") == 12 and vlan_named_in("Users") == 0


def test_vlan_history(local):
    vlans = VlanStore(local)
    domain = vlans.add_domain("Lab")
    vlans.set_vlan(domain.id, 10, "USERS")
    vlans.set_vlan(domain.id, 10, "STAFF", subnets=["10.1.0.0/24"])
    vlans.delete_vlan(domain.id, 10)
    vlans.set_vlan(domain.id, 10, "AGAIN")
    events = vlan_history(local, domain.id, 10)
    assert [event.action for event in events] == ["Added", "Deleted", "Changed", "Added"]
    assert events[2].details == "Name: USERS → STAFF; Subnets: (none) → 10.1.0.0/24"
    assert events[0].subject == "VLAN 10 (AGAIN)"
    assert len(domain_history(local, domain.id)) == 5  # The domain's creation too


def test_name_problems():
    assert name_problem("USERS") == ""
    assert "spaces" in name_problem("Guest Wifi")
    assert "32" in name_problem("x" * 33)


# --------------------------------------------------------------------- Shared through the tribe's server

def tribe_vlans(server, tmp_path, user):
    return TeamVlanStore(team_store(server, tmp_path, user))


def test_vlans_shared_through_the_server(server, tmp_path):
    admin = team_store(server, tmp_path, "admin", ADMIN)
    [network] = admin.import_networks([plan()])
    alice = tribe_vlans(server, tmp_path, "alice")
    domain = alice.add_domain("SIPR", network.id, "CORP")
    alice.set_vlan(domain.id, 10, "USERS", subnets=["10.0.0.0/24"])
    alice.set_vlans(domain.id, [{"vlan": 20, "name": "VOICE"}, {"vlan": 30, "name": "PRINTERS"}])
    bob = tribe_vlans(server, tmp_path, "bob")
    assert [domain.name for domain in bob.domains()] == ["SIPR"]
    assert [(vlan.vlan, vlan.name) for vlan in bob.vlans(domain.id)] == [(10, "USERS"), (20, "VOICE"),
                                                                         (30, "PRINTERS")]
    assert bob.vlan(domain.id, 10).modified_by == "alice (PC)"
    assert bob.vlan(domain.id, 10).subnets == ["10.0.0.0/24"]

    bob.set_vlan(domain.id, 10, "STAFF", subnets=["10.0.0.0/24"])
    with pytest.raises(IpamError, match="changed by bob"):  # Alice hasn't synced since
        alice.set_vlan(domain.id, 10, "ALICE")
    alice.team.sync()
    assert alice.vlan(domain.id, 10).name == "STAFF"
    # Domains are shared too, and changes to the network itself never happen
    assert [subnet.version for subnet in alice.team.subnets(network.id)] == [1]


def test_batch_with_a_conflict_changes_nothing(server, tmp_path):
    admin = team_store(server, tmp_path, "admin", ADMIN)
    [network] = admin.import_networks([plan()])
    alice = tribe_vlans(server, tmp_path, "alice")
    domain = alice.add_domain("SIPR", network.id)
    bob = tribe_vlans(server, tmp_path, "bob")
    alice.set_vlan(domain.id, 30, "ALICE")
    with pytest.raises(IpamError, match="VLAN 30 \\(ALICE\\) was just recorded by alice"):
        bob.set_vlans(domain.id, [{"vlan": 20, "name": "VOICE"}, {"vlan": 30, "name": "BOB"}])
    bob.team.sync()
    assert [vlan.vlan for vlan in bob.vlans(domain.id)] == [30]


def test_offline_vlan_changes_wait_then_reach_the_server(server, tmp_path):
    admin = team_store(server, tmp_path, "admin", ADMIN)
    [network] = admin.import_networks([plan()])
    alice = tribe_vlans(server, tmp_path, "alice")
    domain = alice.add_domain("SIPR", network.id)
    alice.set_vlan(domain.id, 10, "USERS")
    alice.team.close()

    laptop = TeamVlanStore(offline_store(server, tmp_path))
    laptop.set_vlan(domain.id, 20, "VOICE")
    laptop.set_vlan(domain.id, 10, "USERS-2")
    laptop.set_vlan(domain.id, 10, "USERS-3")  # Merged with the change before
    laptop.set_vlan(domain.id, 99, "TEMP")
    laptop.delete_vlan(domain.id, 99)  # Recorded and deleted offline: nothing to send
    assert laptop.vlan(domain.id, 20).name == "VOICE"  # Usable straight away
    assert laptop.pending_count() == 2 and laptop.pending_numbers(domain.id) == {10, 20}
    with pytest.raises(IpamError):
        laptop.set_vlan(domain.id, 30, subnets=["not a subnet"])
    with pytest.raises(ServerUnreachable, match="VLAN domains can't be changed"):
        laptop.add_domain("Offline")
    laptop.team.close()

    laptop = TeamVlanStore(online_again(server, tmp_path))
    assert laptop.pending_count() == 2  # They survive a restart
    assert laptop.flush() == (2, 0)
    bob = tribe_vlans(server, tmp_path, "bob")
    assert [(vlan.vlan, vlan.name) for vlan in bob.vlans(domain.id)] == [(10, "USERS-3"), (20, "VOICE")]


def test_offline_vlan_changes_someone_beat_are_refused(server, tmp_path):
    admin = team_store(server, tmp_path, "admin", ADMIN)
    [network] = admin.import_networks([plan()])
    alice = tribe_vlans(server, tmp_path, "alice")
    domain = alice.add_domain("SIPR", network.id)
    alice.set_vlan(domain.id, 10, "USERS")
    alice.team.close()
    laptop = TeamVlanStore(offline_store(server, tmp_path))
    laptop.set_vlan(domain.id, 20, "ALICE-OFFLINE")
    laptop.set_vlan(domain.id, 10, "ALICE-EDIT")
    laptop.team.close()

    bob = tribe_vlans(server, tmp_path, "bob")
    bob.set_vlan(domain.id, 20, "BOB")
    bob.set_vlan(domain.id, 10, "BOB-EDIT")

    laptop = TeamVlanStore(online_again(server, tmp_path))
    assert laptop.flush() == (0, 2)
    refused = {entry["vlan"]: entry for entry in laptop.refused()}
    assert "was just recorded by bob" in refused[20]["error"] and "changed by bob" in refused[10]["error"]
    assert refused[20]["data"]["name"] == "ALICE-OFFLINE"
    assert laptop.vlan(domain.id, 20) is None and laptop.vlan(domain.id, 10).name == "USERS"
    laptop.discard(refused[10]["seq"])
    assert [entry["vlan"] for entry in laptop.refused()] == [20]


def test_copy_synced_by_an_older_nomad_fetches_vlans_again(server, tmp_path):
    admin = team_store(server, tmp_path, "admin", ADMIN)
    [network] = admin.import_networks([plan()])
    alice = tribe_vlans(server, tmp_path, "alice")
    domain = alice.add_domain("SIPR", network.id)
    alice.set_vlan(domain.id, 10, "USERS")
    # An older NOMAD's copy: synced to the latest revision, without VLAN rows (it didn't keep them)
    path = tmp_path / "old.db"
    old = IpamStore(str(path))
    old.set_meta("revision", alice.team.revision)
    old.set_meta("server_id", key_for(server).server_id)
    old.close()
    from nomad.ipam.client import TeamClient, TeamStore
    key = key_for(server)
    upgraded = TeamStore(key, path, TeamClient(key, user="old", computer="PC"))
    upgraded.sync()
    assert [vlan.name for vlan in TeamVlanStore(upgraded).vlans(domain.id)] == ["USERS"]


def test_older_server_without_vlans(server, tmp_path):
    laptop = team_store(server, tmp_path, "alice")
    laptop.copy.set_meta("server_api", 6)
    vlans = TeamVlanStore(laptop)
    with pytest.raises(IpamError, match="doesn't keep VLANs"):
        vlans.add_domain("SIPR")
    assert vlans.outgoing() == []


# --------------------------------------------------------------------- Bringing in a network map's VLANs

def test_compare_with_map_and_apply(local):
    from nomad.ipam.vlan_compare import ADD, LINK, MISSING, RENAME, apply_vlan_changes, compare_with_map
    from nomad.netmap.vlans import Gateway, MapVlan
    network = local.add_network("Site")
    local.add_subnet(network.id, "10.10.0.0/24", "Users")
    vlans = VlanStore(local)
    domain = vlans.add_domain("CORP", network.id, "CORP")
    vlans.set_vlan(domain.id, 10, "")
    vlans.set_vlan(domain.id, 20, "VOICE-PLAN")
    vlans.set_vlan(domain.id, 50, "OLD")
    vlans.set_vlan(domain.id, 60, "GONE")
    vlans.delete_vlan(domain.id, 60)
    found = [MapVlan(10, "CORP", {"USERS": ["sw1"]}, ["sw1"], gateways=[Gateway("sw1", "10.10.0.1", 24, "Vlan10")]),
             MapVlan(20, "CORP", {"VOICE": ["sw1"]}, ["sw1"]),
             MapVlan(30, "CORP", {"PRINTERS": ["sw1"]}, ["sw1"], gateways=[Gateway("sw1", "10.30.0.1", 24, "Vlan30")]),
             MapVlan(60, "CORP", {"GONE": ["sw1"]}, ["sw1"])]
    changes = compare_with_map(vlans, local, domain, found)
    summary = [(change.action, change.vlan, change.chosen) for change in changes]
    assert summary == [(RENAME, 10, True), (LINK, 10, True), (RENAME, 20, False), (ADD, 30, True), (ADD, 60, False),
                       (MISSING, 50, False)]
    assert "no IPAM subnet holds its VLAN interface 10.30.0.1/24" in changes[3].note
    assert apply_vlan_changes(vlans, domain.id, changes) == 2
    assert (vlans.vlan(domain.id, 10).name, vlans.vlan(domain.id, 10).subnets) == ("USERS", ["10.10.0.0/24"])
    assert vlans.vlan(domain.id, 20).name == "VOICE-PLAN" and vlans.vlan(domain.id, 30).name == "PRINTERS"
    assert vlans.vlan(domain.id, 60) is None
    assert [subnet.cidr for subnet in local.subnets(network.id)] == ["10.10.0.0/24"]  # IPAM untouched


def test_map_interfaces_link_the_ipam_subnet_holding_them(local):
    """A VLAN interface links the most specific IPAM subnet holding its address, whatever either's mask; a domain
    with no network gets the one holding most of them suggested."""
    from nomad.ipam.vlan_compare import LINK, compare_with_map
    from nomad.ipam.vlans import network_for_gateways
    from nomad.netmap.vlans import Gateway, MapVlan
    other = local.add_network("Elsewhere")
    local.add_subnet(other.id, "192.168.9.0/24", "Other")
    network = local.add_network("Site")
    local.add_subnet(network.id, "172.28.0.0/16", "Vlan 6")
    local.add_subnet(network.id, "10.0.2.0/24", "MGMT", "10.0.2.1")
    found = [MapVlan(6, "", {"SIX": ["sw"]}, ["sw"], gateways=[Gateway("sw", "172.28.101.1", 24, "Vl6")]),
             MapVlan(102, "", {"MGMT": ["sw"]}, ["sw"], gateways=[Gateway("sw", "10.0.2.2", 24, "Vl102"),
                                                                  Gateway("r1", "10.0.2.3", 24, "Gi0/0.102")])]
    gateways = [gateway for item in found for gateway in item.gateways]
    assert network_for_gateways(local, gateways).id == network.id
    vlans = VlanStore(local)
    domain = vlans.add_domain("Lab")  # No network: nothing can be linked, and the notes say why
    changes = compare_with_map(vlans, local, domain, found)
    assert all(not change.subnets for change in changes) and "no IPAM network" in changes[0].note
    domain = vlans.update_domain(domain.id, network_id=network.id)
    vlans.set_vlan(domain.id, 6, "SIX")
    changes = {(change.action, change.vlan): change for change in compare_with_map(vlans, local, domain, found)}
    assert changes[(LINK, 6)].subnets == ["172.28.0.0/16"]
    assert changes[("add", 102)].subnets == ["10.0.2.0/24"]  # Two interfaces in it: linked once
