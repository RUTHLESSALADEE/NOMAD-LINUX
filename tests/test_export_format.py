"""The IP Addresses page's exports keep their exact format: Export to Workbook (the tribe's addressing workbook
layout, cell for cell, with its bold headings, column widths and frozen header) and Export to CSV are compared with
saved copies, and nothing the VLANs, Subnet Placement or Network Map pages do changes IPAM's networks, subnets or
addresses, so it never reaches an export.

If an export is meant to change (it should hardly ever be), write the saved copies again (PowerShell) with
    $env:NOMAD_WRITE_EXPORT_SNAPSHOTS = "1"; python -m pytest tests/test_export_format.py; $env:NOMAD_WRITE_EXPORT_SNAPSHOTS = ""
and check the difference in tests/fixtures/export_snapshots.json before committing it.
"""
import csv
import io
import json
import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from test_ipam_server import offline_store, online_again, plan, server, team_store  # noqa: F401 (fixture)

from nomad.ipam.placement import ADVERTISED, IN_PROGRESS, LOCAL, PlacementStore
from nomad.ipam.placement_team import TeamPlacementStore
from nomad.ipam.server import ADMIN
from nomad.ipam.store import RESERVED, USED, IpamStore
from nomad.ipam.vlan_team import TeamVlanStore
from nomad.ipam.vlans import PLANNED, VlanStore
from nomad.ipam.workbook import csv_rows, export_workbook

SNAPSHOTS = Path(__file__).parent / "fixtures" / "export_snapshots.json"
WRITE = os.environ.get("NOMAD_WRITE_EXPORT_SNAPSHOTS") == "1"
WORKBOOK_TABLES = ("networks", "subnets", "addresses")


@pytest.fixture
def local(tmp_path):
    store = IpamStore(str(tmp_path / "ipam.db"), user="tester")
    yield store
    store.close()


def fill(store):
    """Networks with a bit of everything the workbook layout has a place for (and some it doesn't)."""
    sipr = store.add_network("11AB SIPR", "The SIPR page", fields={
        "Unit": "11AB", "Location": "HQ", "Revision date": "2026-04-23", "Revision": "02", "Enclave": "SIPR",
        "ASN": "65139", "UA": "1234", "Imported from": "SIPR Template"})
    subnets = [
        ("10.0.0.0/29", "68900 MAIN Loopback", "", "", {"Telephony Rng": "68900"}, True),
        ("10.0.0.8/29", "68900 MAIN IN-CT", "10.0.0.9", "In-CT", {"Telephony Rng": "68900"}, False),
        ("10.0.0.16/30", "", "", "", {}, False),  # No name: named by its CIDR
        ("10.0.0.20/31", "P2P CORE1-CORE2", "", "", {}, False),
        ("10.0.0.24/32", "RTR1 Lo0", "", "", {}, False),
        ("10.1.0.0/16", "Vlan 7", "10.1.0.1", "", {"Colorless": "Vlan 7"}, False),  # Too big to list every address
        ("10.2.0.0/22", "Block", "", "", {}, False),
        ("10.2.1.0/24", "Inside the block", "10.2.1.1", "", {"Telephony Rng": "68910", "Vlan": "101"}, False),
        ("172.16.5.0/24", "DMVPN Tunnel", "", "Hub and spokes", {}, False),
        ("2001:db8:5::/64", "IPv6 LAN", "2001:db8:5::1", "", {}, False),
    ]
    for cidr, name, gateway, description, fields, loopbacks in subnets:
        store.add_subnet(sipr.id, cidr, name, gateway, description, fields, loopbacks=loopbacks)
    for ip, status, name, mac, description in (
            ("10.0.0.0", RESERVED, "RTR1", "", ""), ("10.0.0.7", USED, "RTR8", "", ""),
            ("10.0.0.9", RESERVED, "GW", "", ""), ("10.0.0.10", USED, "", "00-11-22-33-44-55", "No name: In use"),
            ("10.0.0.16", USED, "on-network", "", ""), ("10.0.0.20", USED, "CORE1", "", ""),
            ("10.0.0.21", USED, "CORE2", "", ""), ("10.0.0.24", USED, "RTR1", "", ""),
            ("10.1.2.3", USED, "big-host", "", "In a /16"), ("10.2.0.5", USED, "in-block", "", ""),
            ("10.2.1.9", RESERVED, "nested-host", "", ""), ("172.16.5.1", USED, "HUB", "", ""),
            ("172.16.5.2", USED, "SPOKE1", "", ""), ("2001:db8:5::10", USED, "v6-host", "", ""),
            ("10.9.9.9", USED, "stray", "", "Outside every subnet")):
        store.set_address(sipr.id, ip, status, name, mac, description)
    plain = store.add_network("Lab: bench/2", fields={"Revision date": "2026-09-01", "ASN": "64512"})
    store.add_subnet(plain.id, "192.168.1.0/24", "Bench LAN", "192.168.1.1")
    store.set_address(plain.id, "192.168.1.20", USED, "printer")
    twin = store.add_network("Lab: bench?2")  # The same page title once its bad characters are swapped
    store.add_subnet(twin.id, "192.168.2.0/25", "Twin")
    return [sipr, plain, twin]


def workbook_snapshot(path):
    workbook = load_workbook(path)
    pages = []
    for sheet in workbook.worksheets:
        pages.append({
            "title": sheet.title,
            "rows": [[value for value in row] for row in sheet.iter_rows(values_only=True)],
            "bold": [cell.coordinate for row in sheet.iter_rows() for cell in row if cell.font and cell.font.bold],
            "widths": {column: sheet.column_dimensions[column].width for column in "ABCDEFGH"},
            "frozen": sheet.freeze_panes,
        })
    return pages


def csv_snapshot(store, network):
    """Export to CSV's text, without when and by whom each line changed (those differ from run to run)."""
    text = io.StringIO()
    writer = csv.writer(text)
    for row in csv_rows(store, network.id, store.subnets(network.id)):
        writer.writerow(row[:-2] if row[0] != "Subnet" else row)
    return list(csv.reader(io.StringIO(text.getvalue())))


def exports(store, networks, tmp_path, name="export"):
    path = tmp_path / f"{name}.xlsx"
    export_workbook(str(path), [(store, store.network(network.id)) for network in networks])
    return {"workbook": workbook_snapshot(path),
            "csv": {network.name: csv_snapshot(store, network) for network in networks}}


def table_rows(store):
    """Every row of IPAM's own tables, deleted ones included: what the exports are made from."""
    return {table: [tuple(row) for row in store.db.execute(f"SELECT * FROM {table} ORDER BY id")]
            for table in WORKBOOK_TABLES}


def snapshot_text(made):
    """The snapshots as JSON with a row on each line, so a change shows up as a readable difference."""
    def rows(items, indent):
        return "[\n" + ",\n".join(indent + json.dumps(item) for item in items) + "\n" + indent[:-1] + "]"

    pages = []
    for page in made["workbook"]:
        fields = [f'  "{key}": ' + (rows(value, "   ") if key == "rows" else json.dumps(value))
                  for key, value in page.items()]
        pages.append(" {\n" + ",\n".join(fields) + "\n }")
    csvs = [f" {json.dumps(name)}: " + rows(lines, "  ") for name, lines in made["csv"].items()]
    return '{"workbook": [\n' + ",\n".join(pages) + '\n],\n"csv": {\n' + ",\n".join(csvs) + "\n}}\n"


def test_exports_match_the_saved_snapshots(local, tmp_path):
    made = exports(local, fill(local), tmp_path)
    made = json.loads(json.dumps(made))  # As saved (tuples become lists)
    if WRITE:
        SNAPSHOTS.write_text(snapshot_text(made), encoding="utf-8")
        pytest.skip("Wrote the export snapshots: check the difference before committing them")
    saved = json.loads(SNAPSHOTS.read_text(encoding="utf-8"))
    for made_page, saved_page in zip(made["workbook"], saved["workbook"]):
        assert made_page == saved_page, f"Export to Workbook changed on page {saved_page['title']}"
    assert len(made["workbook"]) == len(saved["workbook"])
    assert made["csv"] == saved["csv"], "Export to CSV changed"


def test_vlans_and_placement_leave_ipam_and_its_exports_alone(local, tmp_path):
    networks = fill(local)
    sipr = networks[0]
    before_rows, before = table_rows(local), exports(local, networks, tmp_path, "before")

    vlans, placements = VlanStore(local), PlacementStore(local)
    site = vlans.add_domain("SIPR switches", sipr.id, "SIPR-VTP", ranges=[{"first": 100, "last": 199,
                                                                         "name": "Users"}])
    other = vlans.add_domain("Other site", sipr.id)
    vlans.update_domain(site.id, description="Main site")
    vlans.set_vlan(site.id, 7, "MGMT", subnets=["10.1.0.0/16"])
    vlans.set_vlan(site.id, 101, "USERS", subnets=["10.2.1.0/24", "10.0.0.8/29"])
    vlans.set_vlans(site.id, [{"vlan": 102, "name": "VOICE"}, {"vlan": 150, "status": PLANNED}])
    vlans.delete_vlan(site.id, 150)
    vlans.set_vlan(other.id, 30, "OLD", subnets=["192.168.1.0/24"])
    placements.set_placement(sipr.id, "10.1.0.0/16", ADVERTISED, note="one place only")
    placements.set_placement(sipr.id, "10.0.0.16/30", LOCAL, one_segment=True)
    placements.set_placement(sipr.id, "10.0.0.16/30")  # Back to automatic
    placements.set_role(sipr.id, "172.16.5.0/24", "tunnel")
    placements.set_role(sipr.id, "10.0.0.20/31", "other")
    placements.set_role(sipr.id, "10.0.0.20/31")  # Back to automatic
    move = placements.plan_move(sipr.id, "10.2.1.0/24", site.id, 101, "sw1", other.id, 31, "sw2",
                                planned_for="Saturday")
    placements.update_move(move.id, status=IN_PROGRESS)
    placements.complete_move(move.id)
    cancelled = placements.plan_move(sipr.id, "10.1.0.0/16", site.id, 7, "", other.id, 7, "")
    placements.update_move(cancelled.id, status="cancelled")
    vlans.delete_domain(other.id)

    assert table_rows(local) == before_rows
    assert exports(local, networks, tmp_path, "after") == before


def test_tribe_vlans_and_placement_leave_ipam_and_its_exports_alone(server, tmp_path):
    admin = team_store(server, tmp_path, "admin", ADMIN)
    [network] = admin.import_networks([plan()])
    before_rows, before = table_rows(admin.copy), exports(admin, [network], tmp_path, "before")

    alice = team_store(server, tmp_path, "alice")
    vlans, placements = TeamVlanStore(alice), TeamPlacementStore(alice)
    domain = vlans.add_domain("SIPR", network.id, "CORP")
    vlans.set_vlan(domain.id, 10, "USERS", subnets=["10.0.0.0/24"])
    placements.set_placement(network.id, "10.0.0.0/24", LOCAL, note="reused at every site")
    placements.set_role(network.id, "10.0.0.0/24", "routed")
    move =placements.plan_move(network.id, "10.0.0.0/24", from_domain_id=domain.id, from_vlan=10,
                                to_domain_id=domain.id, to_vlan=20)
    placements.complete_move(move.id)
    alice.close()
    offline = TeamVlanStore(offline_store(server, tmp_path))  # Changed offline, sent later
    offline.set_vlan(domain.id, 30, "PRINTERS")
    offline.team.close()
    again = TeamVlanStore(online_again(server, tmp_path))
    assert again.flush() == (1, 0)
    again.team.close()

    admin.sync()
    assert TeamVlanStore(admin).vlan(domain.id, 30).name == "PRINTERS"  # They did reach the server
    assert table_rows(admin.copy) == before_rows
    assert exports(admin, [network], tmp_path, "after") == before
    admin.close()
